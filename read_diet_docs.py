import docx
import os
import sys

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

folder = os.path.join('mitybos_duomenys', 'Mitybos agento informacijos failai')
files = [f for f in os.listdir(folder) if f.endswith('.docx')]

for f in sorted(files):
    if any(k in f for k in ['ANAMNEZ', 'GENERATOR', 'Energij', 'Greiti', 'Meal Prep', 'Sveiki lietuviski']):
        path = os.path.join(folder, f)
        doc = docx.Document(path)
        p_text = '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
        t_text = '\n'.join([c.text.strip() for t in doc.tables for row in t.rows for c in row.cells if c.text.strip()])
        full = p_text + "\n" + t_text
        print(f"==================== {f} (Ilgis: {len(full)}) ====================")
        print(full[:1000])
        print("\n" + "="*50 + "\n")
